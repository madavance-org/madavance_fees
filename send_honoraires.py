import os
import re
import base64
import io
import json
import unicodedata
import requests
import msal
import openpyxl
from datetime import datetime
from urllib.parse import quote

# Configuration
CLIENT_ID = os.environ['CLIENT_ID']
CLIENT_SECRET = os.environ['CLIENT_SECRET']
TENANT_ID = os.environ['TENANT_ID']
SENDER_EMAIL = "domoina@madavance.org"
SHAREPOINT_HOST = "madavancengo.sharepoint.com"
SITE_NAME = "ITMadAvance"
ROOT_FOLDER = "MadAvance/Automatisation frais d'honoraires MadAvance"

MONTHS_FR = {
    1: "janvier", 2: "février", 3: "mars", 4: "avril",
    5: "mai", 6: "juin", 7: "juillet", 8: "août",
    9: "septembre", 10: "octobre", 11: "novembre", 12: "décembre"
}

def get_token():
    app = msal.ConfidentialClientApplication(
        CLIENT_ID,
        authority=f"https://login.microsoftonline.com/{TENANT_ID}",
        client_credential=CLIENT_SECRET
    )
    result = app.acquire_token_for_client(scopes=["https://graph.microsoft.com/.default"])
    if "access_token" not in result:
        raise Exception(f"Échec authentification : {result.get('error_description')}")
    token = result["access_token"]
    print(f"[DEBUG] Token obtenu (débute par) : {token[:30]}...")
    try:
        payload_b64 = token.split(".")[1]
        payload_b64 += "=" * (-len(payload_b64) % 4)
        payload = json.loads(base64.b64decode(payload_b64))
        print(f"[DEBUG] Rôles accordés (roles) : {payload.get('roles', [])}")
        print(f"[DEBUG] App ID (appid) : {payload.get('appid', 'N/A')}")
        print(f"[DEBUG] Tenant (tid) : {payload.get('tid', 'N/A')}")
    except Exception as e:
        print(f"[DEBUG] Impossible de décoder le token : {e}")
    return token

def graph_get(token, url):
    r = requests.get(url, headers={"Authorization": f"Bearer {token}"})
    if not r.ok:
        print(f"[DEBUG] Erreur {r.status_code} sur {url}")
        try:
            print(f"[DEBUG] Détail : {r.json()}")
        except Exception:
            print(f"[DEBUG] Réponse brute : {r.text[:500]}")
        r.raise_for_status()
    return r.json()

def graph_patch(token, url, body):
    r = requests.patch(
        url,
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        json=body
    )
    if not r.ok:
        print(f"[DEBUG] Erreur PATCH {r.status_code} sur {url}")
        try:
            print(f"[DEBUG] Détail : {r.json()}")
        except Exception:
            print(f"[DEBUG] Réponse brute : {r.text[:500]}")
        r.raise_for_status()

def get_drive_id(token):
    site = graph_get(token, f"https://graph.microsoft.com/v1.0/sites/{SHAREPOINT_HOST}:/sites/{SITE_NAME}")
    site_id = site["id"]
    drives = graph_get(token, f"https://graph.microsoft.com/v1.0/sites/{site_id}/drives")
    print("Drives disponibles :")
    for d in drives["value"]:
        print(f"  - {d['name']} ({d['id']})")
    for d in drives["value"]:
        if d["name"].lower() in ("documents", "documents partagés", "shared documents"):
            return d["id"]
    return drives["value"][0]["id"]

def encode_path(path):
    return quote(path, safe='/')

def download_file(token, drive_id, path):
    url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/root:/{encode_path(path)}:/content"
    r = requests.get(url, headers={"Authorization": f"Bearer {token}"})
    r.raise_for_status()
    return r.content

def list_folder(token, drive_id, path):
    url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/root:/{encode_path(path)}:/children"
    return graph_get(token, url)["value"]

def normalize_for_match(name):
    """Normalise un nom de fichier PDF pour la comparaison souple."""
    # Supprimer l'extension .pdf si présente
    if name.lower().endswith(".pdf"):
        name = name[:-4]
    # Normalisation Unicode NFC (é composé vs e + accent)
    name = unicodedata.normalize("NFC", name)
    # Remplacer tirets spéciaux (‐ ‑ – — etc.) par tiret standard
    name = re.sub(r"[‐‑‒–—―﹘﹣－]", "-", name)
    # Supprimer suffixes numériques après l'année : _12, _24, _
    name = re.sub(r"_\d+$", "", name)
    name = re.sub(r"_+$", "", name)
    # Supprimer l'espace parasite après _ : "_ Juin" → "_Juin"
    name = re.sub(r"_\s+", "_", name)
    return name.strip()

def read_excel(content):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    staff = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            entry = dict(zip(headers, row))
            note = entry.get("Notes honoraire")
            if not note or (isinstance(note, str) and note.startswith("=")):
                nom = entry.get("Nom", "")
                mois = entry.get("Mois", "")
                annee = entry.get("Année", "")
                entry["Notes honoraire"] = f"{nom}_{mois}{annee}"
            staff.append(entry)
    return staff

def get_worksheet_name(token, drive_id, item_id):
    url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}/workbook/worksheets"
    result = graph_get(token, url)
    return result["value"][0]["name"]

def setup_status_headers(token, drive_id, item_id, sheet_name):
    """Crée les en-têtes PDF trouvé / Email envoyé en colonnes F et G."""
    url = (
        f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}"
        f"/workbook/worksheets('{quote(sheet_name)}')/range(address='F1:G1')"
    )
    graph_patch(token, url, {"values": [["PDF trouvé", "Email envoyé"]]})

def update_row_status(token, drive_id, item_id, sheet_name, row, pdf_filename, email_address):
    """Met à jour le statut d'une ligne dans l'Excel avec le nom du fichier et l'email destinataire."""
    url = (
        f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}"
        f"/workbook/worksheets('{quote(sheet_name)}')/range(address='F{row}:G{row}')"
    )
    graph_patch(token, url, {"values": [[pdf_filename, email_address]]})

def send_email(token, nom, email, pdf_content, pdf_filename, month, year):
    body = {
        "message": {
            "subject": f"Note d'honoraires – {month.capitalize()} {year}",
            "body": {
                "contentType": "HTML",
                "content": (
                    f"<p>Cher(e) {nom},</p>"
                    f"<p>Veuillez trouver ci-joint votre note d'honoraires pour le mois de {month} {year}.</p>"
                    f"<p>Cordialement,<br>MadAvance</p>"
                )
            },
            "toRecipients": [{"emailAddress": {"address": email, "name": nom}}],
            "attachments": [{
                "@odata.type": "#microsoft.graph.fileAttachment",
                "name": pdf_filename,
                "contentType": "application/pdf",
                "contentBytes": base64.b64encode(pdf_content).decode()
            }]
        }
    }
    r = requests.post(
        f"https://graph.microsoft.com/v1.0/users/{SENDER_EMAIL}/sendMail",
        headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"},
        json=body
    )
    r.raise_for_status()

def main():
    now = datetime.now()
    month = MONTHS_FR[now.month]
    year = now.year
    folder_path = f"{ROOT_FOLDER}/{month} {year}"

    print(f"Traitement du dossier : {folder_path}")

    token = get_token()
    drive_id = get_drive_id(token)
    items = list_folder(token, drive_id, folder_path)

    # Construire la map normalisée : nom_normalisé → nom_réel du fichier
    pdf_map = {}
    for item in items:
        if item["name"].lower().endswith(".pdf"):
            norm = normalize_for_match(item["name"])
            pdf_map[norm] = item["name"]
    print(f"{len(pdf_map)} PDF(s) trouvé(s) dans le dossier")

    # Trouver le fichier Excel
    # Le fichier staff doit contenir "staff" et le nom du mois (ex : staff_septembre _2026.xlsx)
    excel_item = next(
        (i for i in items
         if i["name"].lower().endswith(".xlsx")
         and "staff" in i["name"].lower()
         and month in unicodedata.normalize("NFC", i["name"]).lower()),
        None
    )
    if not excel_item:
        raise Exception(f"Fichier staff du mois de {month} introuvable dans le dossier")

    excel_item_id = excel_item["id"]
    excel_content = download_file(token, drive_id, f"{folder_path}/{excel_item['name']}")
    staff_list = read_excel(excel_content)
    print(f"{len(staff_list)} staff trouvés dans l'Excel")

    # Préparer les colonnes de statut
    sheet_name = get_worksheet_name(token, drive_id, excel_item_id)
    setup_status_headers(token, drive_id, excel_item_id, sheet_name)

    for i, staff in enumerate(staff_list):
        row = i + 2  # ligne 1 = en-têtes, données à partir de la ligne 2
        nom = staff.get("Nom")
        email = staff.get("Email")
        pdf_base = str(staff.get("Notes honoraire"))

        if not nom or not email:
            continue

        # Recherche normalisée du PDF
        pdf_norm = normalize_for_match(pdf_base)
        actual_filename = ""
        email_sent_to = ""

        if pdf_norm not in pdf_map:
            print(f"⚠ PDF introuvable pour {nom} (attendu : {pdf_base})")
        else:
            actual_filename = pdf_map[pdf_norm]
            try:
                pdf_content = download_file(token, drive_id, f"{folder_path}/{actual_filename}")
                send_email(token, nom, email, pdf_content, actual_filename, month, year)
                print(f"✓ Email envoyé à {nom} ({email})")
                email_sent_to = email
            except Exception as e:
                print(f"✗ Erreur pour {nom} : {e}")

        # Écrire le statut dans l'Excel
        update_row_status(token, drive_id, excel_item_id, sheet_name, row, actual_filename, email_sent_to)

if __name__ == "__main__":
    main()
